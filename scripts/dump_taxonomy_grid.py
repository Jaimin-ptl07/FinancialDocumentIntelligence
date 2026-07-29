"""One-time-ish tool: dump the taxonomy xlsx into a committed raw JSON grid so
all downstream parsing lives in pure `fdi.core` (no openpyxl in the backend).

Run with a Python that has openpyxl (NOT the backend venv):
    python scripts/dump_taxonomy_grid.py

Reads:  docs/architecture/taxonomy/Q3VENI-taxonomy-source.xlsx
Writes: docs/architecture/taxonomy/taxonomy.raw.json
"""
import json
import pathlib

import openpyxl

SRC = pathlib.Path("docs/architecture/taxonomy/Q3VENI-taxonomy-source.xlsx")
OUT = pathlib.Path("docs/architecture/taxonomy/taxonomy.raw.json")


def cell(v: object) -> str | None:
    if v is None:
        return None
    return str(v)


def main() -> None:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        grid = [
            [cell(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)
        ]
        while grid and all(x is None for x in grid[-1]):
            grid.pop()
        sheets.append({"title": ws.title, "grid": grid})
    OUT.write_text(
        json.dumps({"schema_version": 1, "sheets": sheets}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    print(f"wrote {OUT}  sheets={len(sheets)} rows={[len(s['grid']) for s in sheets]}")


if __name__ == "__main__":
    main()
